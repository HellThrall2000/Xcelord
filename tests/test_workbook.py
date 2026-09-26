import openpyxl
import pandas as pd
from openpyxl.styles import Font

from xcelord.engine.codegen import parse_reply
from xcelord.excel.context import describe
from xcelord.excel.workbook import Workspace, diff_sheet


def _make_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Sales"])
    ws.append(["a", 10])
    ws.append(["b", 20])
    ws["B2"].font = Font(bold=True)
    ws["D1"] = "=SUM(B2:B3)"  # outside the data frame's columns once pandas reads it? no: it becomes a column
    other = wb.create_sheet("Notes")
    other.append(["Keep", "me"])
    wb.save(path)


def test_diff_sheet_cells():
    old = pd.DataFrame({"A": [1, 2], "B": ["x", None]})
    new = pd.DataFrame({"A": [1, 5], "B": ["x", None], "C": [0, 0]})
    d = diff_sheet(old, new)
    assert d["added_columns"] == ["C"]
    assert [1, 0] in d["cells"] and [0, 0] not in d["cells"]
    assert d["status"] == "modified"


def test_commit_preserves_styles_and_other_sheets(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_book(path)
    ws = Workspace(tmp_path / "work")
    ws.open(path, copy_into_workspace=False)
    assert set(ws.sheets) == {"Data", "Notes"}

    new = ws.snapshot()
    new["Data"].loc[0, "Sales"] = 99
    ws.commit(new)

    wb = openpyxl.load_workbook(path)
    assert wb["Data"]["B2"].value == 99
    assert wb["Data"]["B2"].font.bold  # style kept
    assert wb["Data"]["D1"].value == "=SUM(B2:B3)"  # untouched cell keeps its formula
    assert wb["Notes"]["A1"].value == "Keep"

    assert ws.undo()
    assert openpyxl.load_workbook(path)["Data"]["B2"].value == 10
    assert ws.redo()
    assert openpyxl.load_workbook(path)["Data"]["B2"].value == 99


def test_commit_shape_change_and_new_sheet(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_book(path)
    ws = Workspace(tmp_path / "work")
    ws.open(path, copy_into_workspace=False)
    new = ws.snapshot()
    new["Data"] = new["Data"].iloc[:1]
    new["Summary"] = pd.DataFrame({"Total": [30]})
    del new["Notes"]
    ws.commit(new)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Data", "Summary"]
    assert wb["Data"]["A3"].value is None
    assert wb["Summary"]["A2"].value == 30


def test_edit_cell_coerces(tmp_path):
    path = tmp_path / "book.xlsx"
    _make_book(path)
    ws = Workspace(tmp_path / "work")
    ws.open(path, copy_into_workspace=False)
    ws.edit_cell("Data", 1, 1, "42")
    assert ws.sheets["Data"].iloc[1, 1] == 42


def test_csv_roundtrip(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("A,B\n1,2\n")
    ws = Workspace(tmp_path / "work")
    ws.open(path, copy_into_workspace=True)
    new = ws.snapshot()
    new[ws.active]["C"] = 3
    ws.commit(new)
    assert (tmp_path / "work" / "t.csv").read_text().splitlines()[0] == "A,B,C"


def test_describe_is_compact():
    df = pd.DataFrame({"Dept": ["IT", "HR"] * 50, "Sales": range(100)})
    text = describe({"S": df}, "S")
    assert "'Dept': text" in text and "values: IT, HR" in text
    assert len(text) < 2000


def test_parse_reply_variants():
    assert parse_reply('{"explanation": "Done", "code": "df = df"}') == ("Done", "df = df")
    assert parse_reply('<think>hmm {"x":1}</think>```json\n{"explanation":"E","code":"c"}\n```') == ("E", "c")
    assert parse_reply("Sure!\n```python\ndf['A'] = 1\n```") == ("", "df['A'] = 1")
    assert parse_reply('Here: {"code": "x = {\\"a\\": 1}", "explanation": "ok"} bye')[1] == 'x = {"a": 1}'
