"""In-memory workbook with undo/redo, cell diffs and format-preserving saves."""

from __future__ import annotations

import datetime as dt
import math
import os
import shutil
import tempfile
import threading
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

Sheets = dict[str, pd.DataFrame]

MAX_HISTORY = 30
MAX_ROWS_SENT = 100_000
MAX_DIFF_CELLS = 20_000


# --- value conversion ---------------------------------------------------------

def to_json_value(value):
    if value is None:
        return None
    if isinstance(value, (pd.Timestamp, dt.datetime)):
        if pd.isna(value):
            return None
        if value.hour == value.minute == value.second == 0 and not getattr(value, "microsecond", 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        f = float(value)
        return None if math.isnan(f) or math.isinf(f) else f
    if isinstance(value, pd.Timedelta):
        return str(value)
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value if isinstance(value, (str, int)) else str(value)


def to_excel_value(value):
    if value is None:
        return None
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.to_pydatetime()
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, (str, int, float, bool, dt.datetime, dt.date, dt.time)):
        return value
    return str(value)


def coerce_input(value, series: pd.Series):
    """Convert a value typed into the grid to the column's type where possible."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if pd.api.types.is_bool_dtype(series) and isinstance(value, str):
        if value.strip().lower() in {"true", "yes", "1"}:
            return True
        if value.strip().lower() in {"false", "no", "0"}:
            return False
    if pd.api.types.is_numeric_dtype(series) and not pd.api.types.is_bool_dtype(series):
        try:
            num = float(value)
            return int(num) if num.is_integer() and pd.api.types.is_integer_dtype(series) else num
        except (TypeError, ValueError):
            return value
    if pd.api.types.is_datetime64_any_dtype(series):
        parsed = pd.to_datetime(value, errors="coerce")
        return value if pd.isna(parsed) else parsed
    return value


# --- loading ------------------------------------------------------------------

def _parse_dates(df: pd.DataFrame) -> pd.DataFrame:
    for col in df.columns:
        s = df[col]
        if not (pd.api.types.is_object_dtype(s) or pd.api.types.is_string_dtype(s)):
            continue
        non_null = s.dropna()
        if non_null.empty:
            continue
        sample = non_null.astype(str).head(200)
        if not sample.str.contains(r"\d{1,4}[-/.]\d{1,2}", regex=True).mean() > 0.8:
            continue
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            parsed = pd.to_datetime(s, errors="coerce", format="mixed")
        if parsed.notna().sum() / len(non_null) > 0.9:
            df[col] = parsed
    return df


def read_file(path: Path) -> Sheets:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        sheets = {path.stem[:31] or "Sheet1": pd.read_csv(path)}
    elif suffix in {".xlsx", ".xlsm"}:
        sheets = pd.read_excel(path, sheet_name=None, engine="openpyxl")
        if not sheets:
            sheets = {"Sheet1": pd.DataFrame()}
    else:
        raise ValueError("Only .xlsx, .xlsm and .csv files are supported")
    for name, df in sheets.items():
        df.columns = [str(c) for c in df.columns]
        sheets[name] = _parse_dates(df)
    return sheets


# --- diffing ------------------------------------------------------------------

def _neq(a: pd.Series, b: pd.Series) -> np.ndarray:
    a = a.reset_index(drop=True)
    b = b.reset_index(drop=True)
    both_na = (a.isna() & b.isna()).to_numpy()
    try:
        eq = (a == b).fillna(False).to_numpy(dtype=bool)
    except (TypeError, ValueError):
        eq = (a.astype(str) == b.astype(str)).to_numpy(dtype=bool)
    return ~(eq | both_na)


def diff_sheet(old: pd.DataFrame | None, new: pd.DataFrame) -> dict:
    """Cells that differ, keyed by the *new* frame's positions."""
    if old is None:
        return {"status": "added", "cells": [], "added_columns": list(new.columns),
                "removed_columns": [], "rows_before": 0, "rows_after": len(new), "changed": len(new) * len(new.columns)}
    added = [c for c in new.columns if c not in old.columns]
    removed = [c for c in old.columns if c not in new.columns]
    cells: list[list[int]] = []
    changed = 0
    same_rows = len(old) == len(new)
    col_pos = {c: i for i, c in enumerate(new.columns)}
    for col in new.columns:
        j = col_pos[col]
        if col in added:
            changed += len(new)
            if len(cells) < MAX_DIFF_CELLS:
                cells.extend([i, j] for i in range(min(len(new), MAX_DIFF_CELLS - len(cells))))
            continue
        if not same_rows:
            continue
        rows = np.flatnonzero(_neq(old[col], new[col]))
        changed += len(rows)
        for i in rows[: max(0, MAX_DIFF_CELLS - len(cells))]:
            cells.append([int(i), j])
    status = "unchanged"
    if changed or added or removed or not same_rows or list(old.columns) != list(new.columns):
        status = "modified"
    return {
        "status": status,
        "cells": cells,
        "added_columns": added,
        "removed_columns": removed,
        "rows_before": len(old),
        "rows_after": len(new),
        "changed": changed,
    }


def diff_workbook(old: Sheets, new: Sheets) -> dict[str, dict]:
    out = {name: diff_sheet(old.get(name), df) for name, df in new.items()}
    for name in old:
        if name not in new:
            out[name] = {"status": "removed", "cells": [], "added_columns": [], "removed_columns": [],
                         "rows_before": len(old[name]), "rows_after": 0, "changed": 0}
    return out


def sheet_payload(df: pd.DataFrame) -> dict:
    view = df.head(MAX_ROWS_SENT)
    rows = [[to_json_value(v) for v in row] for row in view.itertuples(index=False, name=None)]
    types = []
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_bool_dtype(s):
            types.append("bool")
        elif pd.api.types.is_numeric_dtype(s):
            types.append("number")
        elif pd.api.types.is_datetime64_any_dtype(s):
            types.append("date")
        else:
            types.append("text")
    return {"columns": [str(c) for c in df.columns], "types": types, "rows": rows,
            "total_rows": len(df), "truncated": len(df) > MAX_ROWS_SENT}


# --- saving -------------------------------------------------------------------

def _write_full(ws, df: pd.DataFrame) -> None:
    old_rows, old_cols = ws.max_row, ws.max_column
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=1, column=j, value=str(col))
    for i, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=to_excel_value(value))
    new_rows, new_cols = len(df) + 1, len(df.columns)
    for i in range(1, old_rows + 1):
        for j in range(1, old_cols + 1):
            if i > new_rows or j > new_cols:
                ws.cell(row=i, column=j).value = None


def write_file(path: Path, sheets: Sheets, previous: Sheets | None) -> None:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        next(iter(sheets.values())).to_csv(path, index=False)
        return

    wb = openpyxl.load_workbook(path, keep_vba=suffix == ".xlsm") if path.exists() else openpyxl.Workbook()
    fresh = not path.exists()
    for name in list(wb.sheetnames):
        if name not in sheets and (not fresh or len(wb.sheetnames) > 1):
            del wb[name]
    for name, df in sheets.items():
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            _write_full(ws, df)
            continue
        ws = wb[name]
        old = (previous or {}).get(name)
        if old is not None and list(old.columns) == list(df.columns) and len(old) == len(df):
            # Same shape: touch only changed cells so formulas/styles elsewhere survive.
            for j, col in enumerate(df.columns):
                for i in np.flatnonzero(_neq(old[col], df[col])):
                    ws.cell(row=int(i) + 2, column=j + 1, value=to_excel_value(df[col].iloc[int(i)]))
        elif old is None or not old.equals(df):
            _write_full(ws, df)
    if fresh and "Sheet" in wb.sheetnames and "Sheet" not in sheets and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    # Keep workbook order aligned with the in-memory order.
    wb._sheets.sort(key=lambda ws: list(sheets).index(ws.title) if ws.title in sheets else len(sheets))

    fd, tmp = tempfile.mkstemp(suffix=path.suffix, dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


# --- workspace ----------------------------------------------------------------

def _copy(sheets: Sheets) -> Sheets:
    return {name: df.copy() for name, df in sheets.items()}


class Workspace:
    """The single open workbook. All mutations go through commit() for undo."""

    def __init__(self, workdir: Path):
        self.workdir = workdir
        self.workdir.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.path: Path | None = None
        self.sheets: Sheets = {}
        self.active: str | None = None
        self._undo: list[tuple[Sheets, str | None]] = []
        self._redo: list[tuple[Sheets, str | None]] = []

    @property
    def loaded(self) -> bool:
        return self.path is not None

    def open(self, path: Path, copy_into_workspace: bool) -> None:
        path = path.expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"{path} does not exist")
        if copy_into_workspace:
            target = self.workdir / path.name
            if target != path:
                shutil.copyfile(path, target)
            path = target
        sheets = read_file(path)
        with self.lock:
            self.path = path
            self.sheets = sheets
            self.active = next(iter(sheets))
            self._undo.clear()
            self._redo.clear()

    def state(self) -> dict:
        with self.lock:
            if not self.loaded:
                return {"loaded": False}
            return {
                "loaded": True,
                "name": self.path.name,
                "path": str(self.path),
                "sheets": [{"name": n, "rows": len(df), "cols": len(df.columns)} for n, df in self.sheets.items()],
                "active": self.active,
                "can_undo": bool(self._undo),
                "can_redo": bool(self._redo),
            }

    def snapshot(self) -> Sheets:
        with self.lock:
            return _copy(self.sheets)

    def set_active(self, name: str) -> None:
        with self.lock:
            if name not in self.sheets:
                raise KeyError(name)
            self.active = name

    def commit(self, new_sheets: Sheets, active: str | None = None) -> None:
        with self.lock:
            previous = self.sheets
            write_file(self.path, new_sheets, previous)
            self._undo.append((previous, self.active))
            del self._undo[:-MAX_HISTORY]
            self._redo.clear()
            self.sheets = new_sheets
            if active and active in new_sheets:
                self.active = active
            elif self.active not in new_sheets:
                self.active = next(iter(new_sheets))

    def _swap(self, source: list, target: list) -> bool:
        with self.lock:
            if not source:
                return False
            sheets, active = source.pop()
            write_file(self.path, sheets, self.sheets)
            target.append((self.sheets, self.active))
            self.sheets = sheets
            self.active = active if active in sheets else next(iter(sheets))
            return True

    def undo(self) -> bool:
        return self._swap(self._undo, self._redo)

    def redo(self) -> bool:
        return self._swap(self._redo, self._undo)

    def edit_cell(self, sheet: str, row: int, col: int, value) -> None:
        with self.lock:
            df = self.sheets[sheet]
            if not (0 <= row < len(df) and 0 <= col < len(df.columns)):
                raise IndexError("Cell out of range")
            new = _copy(self.sheets)
            frame = new[sheet]
            column = frame.columns[col]
            coerced = coerce_input(value, frame[column])
            try:
                frame.iloc[row, col] = coerced
            except (TypeError, ValueError):
                frame[column] = frame[column].astype(object)
                frame.iloc[row, col] = coerced
            self.commit(new)
